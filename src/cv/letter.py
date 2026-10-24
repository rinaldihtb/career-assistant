"""Surat lamaran dari profil CV dan satu lowongan, dengan tepat 1 panggilan LLM.

LLM hanya menulis paragraf isi (structured output). Salam pembuka, salam penutup, dan nama
penanda tangan disusun Python, sehingga nama tidak dikirim ke LLM dan tidak bisa salah tulis.

Prompt berisi profil (posisi tujuan, pendidikan, skill, pengalaman per jabatan) dan lowongan
(judul, perusahaan, lokasi, tipe kerja, skill cocok dan skill kurang hasil ``match_jobs``,
deskripsi). Teks CV, nama, dan lokasi pelamar tidak dikirim. Total lama pengalaman juga tidak
dikirim, agar surat tidak menyebut jumlah tahun yang tidak tertulis di CV.

LLM disuntikkan, fungsi tidak pernah melempar exception, dan kegagalan ada di
``CoverLetterResult.error``. Halaman surat meng-cache hasil per ``version`` profil dan lowongan.
"""

import io
import re
from typing import Literal

from docx import Document
from langchain_core.language_models import BaseChatModel
from langchain_core.messages import HumanMessage, SystemMessage
from pydantic import BaseModel, Field

from agents.prompts import CV_LETTER_SYSTEM, CV_LETTER_USER
from cv.match import JobMatch
from cv.profile import CVProfile, Experience
from models.agent import LLMUsage

LetterLanguage = Literal["id", "en"]
LANGUAGE_NAMES: dict[str, str] = {"id": "Bahasa Indonesia", "en": "English"}
LETTER_LENGTHS = (150, 250, 350)
NAME_PLACEHOLDER = "[Nama Anda]"
# 16 dari 473 deskripsi lebih panjang dari ini (terpanjang 9.407 karakter). Bagian awal
# deskripsi sudah memuat posisi, tanggung jawab, dan kualifikasi.
MAX_JOB_CHARS = 4000
# Nama pengganti dari situs lowongan, bukan nama perusahaan: 18 dari 473 lowongan. Tanpa ini
# surat dibuka dengan "Yth. Tim Rekrutmen Pengiklan Anonim,".
ANONYMOUS_COMPANIES = frozenset({"pengiklan anonim"})
GREETINGS = {"id": "Yth. Tim Rekrutmen {company},", "en": "Dear {company} Hiring Team,"}
GREETINGS_NO_COMPANY = {"id": "Yth. Tim Rekrutmen,", "en": "Dear Hiring Team,"}
CLOSINGS = {"id": "Hormat saya,", "en": "Sincerely,"}
_FILE_NAME_UNSAFE = re.compile(r"[^\w\- ]+")


class LetterSettings(BaseModel):
    """Pilihan pengguna untuk surat: bahasa dan target panjang (kata)."""

    model_config = {"frozen": True}

    language: LetterLanguage = "id"
    words: int = 250


class CoverLetter(BaseModel):
    """Skema structured output. Satu field wajib tanpa default (mode strict)."""

    paragraphs: list[str] = Field(
        description="3-4 paragraf isi surat, tanpa salam pembuka, salam penutup, atau nama."
    )


class CoverLetterResult(BaseModel):
    """Hasil ``generate_cover_letter``: ``letter`` terisi jika berhasil, selain itu ``error``."""

    letter: CoverLetter | None = None
    usage: LLMUsage = Field(default_factory=LLMUsage)
    error: str | None = None


def letter_prompt_input(
    profile: CVProfile, match: JobMatch, description: str | None
) -> tuple[str, str]:
    """Teks ``{profile}`` dan ``{job}`` untuk ``CV_LETTER_USER``.

    Returns:
        ``(profil, lowongan)``. Baris yang datanya kosong dilewati; deskripsi dirapikan
        (spasi dan baris kosong) lalu dipotong ``MAX_JOB_CHARS``.
    """
    education = " ".join(p for p in (profile.education_level, profile.education_major) if p)
    experiences = [_experience_line(exp) for exp in profile.experiences if exp.title.strip()]
    profile_lines = [
        f"Posisi tujuan: {profile.target_role}" if profile.target_role else "",
        f"Pendidikan: {education}" if education else "",
        f"Skill: {', '.join(profile.skills)}" if profile.skills else "",
        "Pengalaman:\n" + "\n".join(experiences) if experiences else "Pengalaman: tidak ada",
    ]
    job_lines = [
        f"Posisi: {match.job_title}",
        f"Perusahaan: {company}" if (company := hiring_company(match.company_name)) else "",
        f"Lokasi: {' '.join(match.location.split())}" if match.location else "",
        f"Tipe kerja: {match.work_type}" if match.work_type else "",
        f"Skill cocok dengan profil: {', '.join(match.matched_skills)}"
        if match.matched_skills
        else "",
        f"Skill diminta, tidak ada di profil: {', '.join(match.missing_skills)}"
        if match.missing_skills
        else "",
        f"Deskripsi:\n{_clip_description(description)}" if description else "",
    ]
    profile_text = "\n".join(line for line in profile_lines if line)
    job_text = "\n".join(line for line in job_lines if line)
    return _escape_tag(profile_text, "profil"), _escape_tag(job_text, "lowongan")


def generate_cover_letter(
    profile: CVProfile,
    match: JobMatch,
    description: str | None,
    *,
    llm: BaseChatModel,
    settings: LetterSettings | None = None,
) -> CoverLetterResult:
    """Tulis paragraf isi surat lamaran dengan tepat satu panggilan LLM.

    Args:
        profile: Profil CV aktif (mungkin sudah dikoreksi pengguna).
        match: Lowongan tujuan dari hasil ``match_jobs``.
        description: Deskripsi lowongan dari database; ``None`` jika tidak tersedia.
        llm: Chat model; dipanggil lewat ``with_structured_output(CoverLetter)``.
        settings: Bahasa dan target panjang; default Bahasa Indonesia, 250 kata.

    Returns:
        ``CoverLetterResult``. Paragraf kosong dibuang; tanpa paragraf tersisa dianggap
        gagal. Tidak pernah melempar exception.
    """
    settings = settings or LetterSettings()
    usage = LLMUsage()
    profile_text, job_text = letter_prompt_input(profile, match, description)
    messages = [
        SystemMessage(CV_LETTER_SYSTEM),
        HumanMessage(
            CV_LETTER_USER.format(
                profile=profile_text,
                job=job_text,
                language=LANGUAGE_NAMES[settings.language],
                words=settings.words,
            )
        ),
    ]
    writer = llm.with_structured_output(CoverLetter, include_raw=True)
    try:
        output = writer.invoke(messages)
    except Exception as exc:  # noqa: BLE001 - tipe error provider LLM sangat beragam
        return CoverLetterResult(usage=usage, error=f"LLM error: {exc}")

    usage.add(output["raw"], step="cv_letter")
    parsed: CoverLetter | None = output["parsed"]
    if parsed is None:
        return CoverLetterResult(
            usage=usage, error=f"Output LLM tidak valid: {output['parsing_error']}"
        )
    paragraphs = [" ".join(p.split()) for p in parsed.paragraphs if p.strip()]
    if not paragraphs:
        return CoverLetterResult(usage=usage, error="Surat kosong")
    return CoverLetterResult(letter=CoverLetter(paragraphs=paragraphs), usage=usage)


def signer_name(full_name: str | None) -> str:
    """Nama penanda tangan dari profil; ``NAME_PLACEHOLDER`` jika kosong.

    Nama yang seluruhnya huruf kapital atau huruf kecil (3 dari 5 CV fiktif set eval ditulis
    kapital semua) diubah menjadi huruf awal kapital per kata: "SITI NUR'AINI" menjadi
    "Siti Nur'aini", "S.KOM." menjadi "S.Kom.". Nama dengan huruf campuran dibiarkan, karena
    penulisannya sudah dipilih pemilik CV ("McDonald", "dr. Rina").
    """
    name = " ".join((full_name or "").split())
    if not name:
        return NAME_PLACEHOLDER
    if not (name.isupper() or name.islower()):
        return name
    return " ".join("-".join(_capitalize_part(p) for p in word.split("-")) for word in name.split())


def letter_to_text(
    letter: CoverLetter,
    *,
    company: str | None,
    name: str,
    language: LetterLanguage = "id",
) -> str:
    """Surat lengkap sebagai teks: salam pembuka, paragraf isi, salam penutup, dan nama.

    Antarbagian dipisah satu baris kosong; salam penutup dan nama dipisah satu baris baru.
    """
    company = hiring_company(company)
    greeting = (
        GREETINGS[language].format(company=company) if company else GREETINGS_NO_COMPANY[language]
    )
    blocks = [greeting, *letter.paragraphs, f"{CLOSINGS[language]}\n{name}"]
    return "\n\n".join(blocks)


def letter_to_docx(text: str, *, author: str = "") -> bytes:
    """File .docx dari teks surat (mungkin sudah diedit pengguna). Tanpa LLM.

    Blok yang dipisah baris kosong menjadi paragraf Word; baris baru di dalam blok menjadi
    line break (misalnya antara salam penutup dan nama).
    """
    document = Document()
    document.core_properties.author = author
    document.core_properties.title = "Surat lamaran"
    for block in re.split(r"\n\s*\n", text.strip()):
        paragraph = document.add_paragraph()
        for index, line in enumerate(block.splitlines()):
            if index:
                paragraph.add_run().add_break()
            paragraph.add_run(line.strip())
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def letter_file_name(match: JobMatch) -> str:
    """Nama file unduhan, misalnya "Surat lamaran - Data Analyst - PT Contoh.docx"."""
    parts = ["Surat lamaran", match.job_title, hiring_company(match.company_name) or ""]
    cleaned = [" ".join(_FILE_NAME_UNSAFE.sub(" ", part).split()) for part in parts]
    return " - ".join(part for part in cleaned if part)[:120] + ".docx"


def hiring_company(name: str | None) -> str | None:
    """Nama perusahaan untuk surat; ``None`` jika kosong atau nama pengganti (Pengiklan Anonim)."""
    name = " ".join((name or "").split())
    if not name or name.casefold() in ANONYMOUS_COMPANIES:
        return None
    return name


def word_count(text: str) -> int:
    """Jumlah kata (dipisah spasi), untuk caption panjang surat."""
    return len(text.split())


def _experience_line(exp: Experience) -> str:
    """Baris pengalaman: jabatan, perusahaan, dan periode jika tertulis."""
    title = " ".join(exp.title.split())
    company = f", {' '.join(exp.company.split())}" if exp.company else ""
    end = "sekarang" if (exp.end or "").strip().casefold() == "present" else exp.end
    period = f" ({exp.start} s.d. {end})" if exp.start and end else ""
    return f"- {title}{company}{period}"


def _clip_description(text: str | None) -> str:
    """Rapikan spasi per baris, buang baris kosong, lalu potong ``MAX_JOB_CHARS``."""
    lines = [" ".join(line.split()) for line in (text or "").splitlines()]
    cleaned = "\n".join(line for line in lines if line)
    if len(cleaned) <= MAX_JOB_CHARS:
        return cleaned
    return cleaned[:MAX_JOB_CHARS].rsplit(" ", 1)[0] + " ..."


def _escape_tag(text: str, tag: str) -> str:
    """Cegah data menutup tag prompt lebih awal (pola sama dengan ``</job>`` di RAG agent)."""
    return text.replace(f"</{tag}>", f"</ {tag}>")


def _capitalize_part(part: str) -> str:
    """Huruf awal kapital per segmen titik: "S.KOM." -> "S.Kom."; "NUR'AINI" -> "Nur'aini"."""
    return ".".join(segment.capitalize() for segment in part.split("."))
