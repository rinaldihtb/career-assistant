"""File surat lamaran (.docx dan .pdf) dari teks surat yang mungkin sudah diedit. Tanpa LLM.

Kedua format dirender dari satu daftar blok (``letter_blocks``) dan satu tabel gaya
(``STYLES``), sehingga tata letaknya sama: kop (nama besar di tengah dan baris kontak bergaris
bawah), tanggal, penerima, perihal tebal, isi rata kiri, salam penutup, ruang tanda tangan,
dan nama tebal. A4, margin 2,3 cm kiri-kanan dan 2,0 cm atas-bawah, huruf 11 pt, spasi 1,15.

Font .docx Cambria (tersedia di Word); PDF menanamkan Caladea (lisensi OFL, ukuran huruf sama
dengan Cambria) dari ``assets/fonts``, karena Cambria tidak boleh disebarkan bersama aplikasi.
"""

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from fpdf import FPDF

Role = Literal["name", "contact", "date", "header", "subject", "body", "closing", "signer"]

PAGE_WIDTH_CM, PAGE_HEIGHT_CM = 21.0, 29.7
MARGIN_SIDE_CM, MARGIN_TOP_CM = 2.3, 2.0
FONT_SIZE = 11
LINE_SPACING = 1.15
DOCX_FONT = "Cambria"
PDF_FONT = "Caladea"
FONT_DIR = Path(__file__).resolve().parents[2] / "assets" / "fonts"
# Tinggi baris bawaan Cambria/Caladea sekitar 1,17 x ukuran huruf; PDF tidak punya spasi
# "1,15 kali" seperti Word, jadi tinggi baris dihitung ukuran x 1,17 x 1,15.
FONT_LINE_HEIGHT = 1.17
# Baris yang diawali kata ini (tanpa beda huruf besar) adalah perihal surat.
_SUBJECT = re.compile(r"^\s*(?:perihal|hal|re|subject)\s*:", re.IGNORECASE)


@dataclass(frozen=True)
class BlockStyle:
    """Gaya satu jenis blok, dipakai renderer .docx dan .pdf."""

    size: float = FONT_SIZE
    bold: bool = False
    center: bool = False
    before: float = 0  # pt
    after: float = 9  # pt
    letter_spacing: float = 0  # pt
    rule_below: bool = False


STYLES: dict[str, BlockStyle] = {
    "name": BlockStyle(size=17, bold=True, center=True, after=0, letter_spacing=1.5),
    "contact": BlockStyle(size=8.5, center=True, before=3, after=2, rule_below=True),
    "date": BlockStyle(before=20, after=14),
    "header": BlockStyle(after=14),
    "subject": BlockStyle(bold=True, after=12),
    "body": BlockStyle(),
    "closing": BlockStyle(before=5, after=30),
    "signer": BlockStyle(bold=True, after=0),
}


@dataclass(frozen=True)
class LetterBlock:
    """Satu paragraf surat: jenisnya dan baris-barisnya (dipisah line break, bukan paragraf)."""

    role: Role
    lines: tuple[str, ...]


def letter_blocks(text: str, *, contact: str = "", fallback_name: str = "") -> list[LetterBlock]:
    """Pecah teks surat menjadi blok berjenis.

    - Blok dipisah baris kosong. Blok terakhir yang diawali salam berakhiran koma ("Hormat
      saya,") lalu nama dipecah menjadi ``closing`` dan ``signer``.
    - Kop: nama penanda tangan (dari teks, supaya mengikuti editan pengguna; jika tidak ada,
      ``fallback_name``) dalam huruf kapital, lalu ``contact`` jika diisi.
    - Jika ada blok perihal ("Perihal:", "Re:"), blok pertama sebelumnya adalah tanggal dan
      sisanya penerima. Tanpa perihal, semua blok sebelum penutup adalah isi.
    """
    blocks = [
        tuple(line.strip() for line in block.splitlines())
        for block in re.split(r"\n\s*\n", text.strip())
        if block.strip()
    ]
    closing: list[LetterBlock] = []
    if len(blocks) > 1 and len(blocks[-1]) > 1 and blocks[-1][0].endswith(","):
        last = blocks.pop()
        closing = [LetterBlock("closing", last[:1]), LetterBlock("signer", last[1:])]

    signer = closing[1].lines[0] if closing else fallback_name.strip()
    kop = [LetterBlock("name", (signer.upper(),))] if signer else []
    if contact.strip():
        kop.append(LetterBlock("contact", (contact.strip(),)))

    subject = next((i for i, lines in enumerate(blocks) if _SUBJECT.match(lines[0])), None)
    body = []
    for index, lines in enumerate(blocks):
        if subject is None or index > subject:
            role: Role = "body"
        elif index == subject:
            role = "subject"
        else:
            role = "date" if index == 0 else "header"
        body.append(LetterBlock(role, lines))
    return [*kop, *body, *closing]


def letter_to_docx(
    text: str, *, author: str = "", contact: str = "", title: str = "Surat lamaran"
) -> bytes:
    """File .docx dari teks surat. Lihat ``letter_blocks`` untuk struktur dan ``STYLES``."""
    document = Document()
    document.core_properties.author = author
    document.core_properties.title = title
    for section in document.sections:
        section.page_width, section.page_height = Cm(PAGE_WIDTH_CM), Cm(PAGE_HEIGHT_CM)
        section.left_margin = section.right_margin = Cm(MARGIN_SIDE_CM)
        section.top_margin = section.bottom_margin = Cm(MARGIN_TOP_CM)
    normal = document.styles["Normal"]
    normal.font.name = DOCX_FONT
    normal.font.size = Pt(FONT_SIZE)
    # Font Asia Timur ikut diisi; tanpa itu Word memakai font tema di sebagian teks.
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), DOCX_FONT)
    normal.paragraph_format.line_spacing = LINE_SPACING

    for block in letter_blocks(text, contact=contact, fallback_name=author):
        style = STYLES[block.role]
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if style.center else WD_ALIGN_PARAGRAPH.LEFT
        )
        paragraph.paragraph_format.space_before = Pt(style.before)
        paragraph.paragraph_format.space_after = Pt(style.after)
        if block.role in ("name", "contact"):
            paragraph.paragraph_format.line_spacing = 1.0
        for index, line in enumerate(block.lines):
            if index:
                paragraph.add_run().add_break()
            run = paragraph.add_run(line)
            run.bold = style.bold
            run.font.size = Pt(style.size)
            if style.letter_spacing:
                _set_letter_spacing(run, style.letter_spacing)
        if style.rule_below:
            _add_bottom_border(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def letter_to_pdf(
    text: str, *, author: str = "", contact: str = "", title: str = "Surat lamaran"
) -> bytes:
    """File .pdf dari teks surat dengan tata letak yang sama dengan ``letter_to_docx``.

    Satuan pt; teks dibungkus otomatis selebar area tulis. Font Caladea ditanamkan (subset).
    """
    pdf = FPDF(unit="pt", format="A4")
    side, top = _cm_to_pt(MARGIN_SIDE_CM), _cm_to_pt(MARGIN_TOP_CM)
    pdf.set_margins(side, top, side)
    pdf.set_auto_page_break(True, margin=top)
    pdf.add_font(PDF_FONT, "", str(FONT_DIR / "Caladea-Regular.ttf"))
    pdf.add_font(PDF_FONT, "B", str(FONT_DIR / "Caladea-Bold.ttf"))
    pdf.set_author(author)
    pdf.set_title(title)
    pdf.add_page()

    for index, block in enumerate(letter_blocks(text, contact=contact, fallback_name=author)):
        style = STYLES[block.role]
        if index:
            pdf.set_y(pdf.get_y() + style.before)
        pdf.set_font(PDF_FONT, "B" if style.bold else "", style.size)
        pdf.set_char_spacing(style.letter_spacing)
        spacing = 1.0 if block.role in ("name", "contact") else LINE_SPACING
        pdf.multi_cell(
            0,
            style.size * FONT_LINE_HEIGHT * spacing,
            "\n".join(block.lines),
            align="C" if style.center else "L",
            new_x="LMARGIN",
            new_y="NEXT",
        )
        if style.rule_below:
            y = pdf.get_y() + 1
            pdf.set_line_width(0.5)
            pdf.line(side, y, pdf.w - side, y)
            pdf.set_y(y)
        pdf.set_y(pdf.get_y() + style.after)
    pdf.set_char_spacing(0)
    return bytes(pdf.output())


def _cm_to_pt(cm: float) -> float:
    return cm / 2.54 * 72


def _set_letter_spacing(run, points: float) -> None:
    """Jarak antarhuruf (w:spacing, satuan 1/20 pt) untuk nama di kop."""
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(round(points * 20)))
    run._r.get_or_add_rPr().append(spacing)


def _add_bottom_border(paragraph) -> None:
    """Garis bawah tipis di bawah baris kontak (w:pBdr/w:bottom)."""
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "000000")):
        bottom.set(qn(f"w:{key}"), value)
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)
