"""File surat (.docx dan .pdf) dari teks surat: struktur blok, format, dan isi yang sama."""

import io

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pypdf import PdfReader

from src.cv.letter_file import STYLES, LetterBlock, letter_blocks, letter_to_docx, letter_to_pdf

CONTACT = "Malang  ·  0812-3456-7890  ·  nadia@mail.com"
LETTER = (
    "Malang, 5 Oktober 2026\n\n"
    "Kepada Yth.\nTim Rekrutmen\nPT Contoh\n\n"
    "Perihal: Lamaran Pekerjaan sebagai Data Analyst\n\n"
    "Dengan hormat,\n\n"
    "Isi pertama.\n\n"
    "Isi kedua.\n\n"
    "Hormat saya,\nNadia Contoh"
)


def _roles(blocks):
    return [(block.role, "\n".join(block.lines)) for block in blocks]


def test_letter_blocks_split_kop_header_body_and_closing():
    assert _roles(letter_blocks(LETTER, contact=CONTACT)) == [
        ("name", "NADIA CONTOH"),
        ("contact", CONTACT),
        ("date", "Malang, 5 Oktober 2026"),
        ("header", "Kepada Yth.\nTim Rekrutmen\nPT Contoh"),
        ("subject", "Perihal: Lamaran Pekerjaan sebagai Data Analyst"),
        ("body", "Dengan hormat,"),
        ("body", "Isi pertama."),
        ("body", "Isi kedua."),
        ("closing", "Hormat saya,"),
        ("signer", "Nadia Contoh"),
    ]


def test_kop_follows_edited_signer_and_skips_empty_contact():
    blocks = letter_blocks("Isi.\n\nSincerely,\nNama Baru", fallback_name="Nama Lama")

    assert _roles(blocks) == [
        ("name", "NAMA BARU"),
        ("body", "Isi."),
        ("closing", "Sincerely,"),
        ("signer", "Nama Baru"),
    ]


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        # Tanpa perihal: semua blok sebelum penutup adalah isi.
        ("Tanggal\n\nIsi.", [("name", "N"), ("body", "Tanggal"), ("body", "Isi.")]),
        # Tanpa penutup: kop memakai nama cadangan, blok terakhir tetap isi.
        (
            "Re: X\n\nBaris satu\nBaris dua",
            [("name", "N"), ("subject", "Re: X"), ("body", "Baris satu\nBaris dua")],
        ),
        ("", [("name", "N")]),
    ],
)
def test_letter_blocks_without_subject_or_closing(text, expected):
    assert _roles(letter_blocks(text, fallback_name="N")) == expected


def test_docx_follows_layout():
    document = Document(io.BytesIO(letter_to_docx(LETTER, author="Nadia", contact=CONTACT)))
    section = document.sections[0]
    name, contact, date_, *_, subject = document.paragraphs[:5]
    signer = document.paragraphs[-1]

    assert (round(section.page_width.cm, 1), round(section.page_height.cm, 1)) == (21.0, 29.7)
    assert round(section.left_margin.cm, 1) == 2.3
    assert round(section.top_margin.cm, 1) == 2.0
    assert document.styles["Normal"].font.name == "Cambria"
    assert document.styles["Normal"].font.size.pt == 11
    assert name.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert name.runs[0].bold and name.runs[0].font.size.pt == 17
    assert contact.runs[0].font.size.pt == 8.5
    assert (
        contact._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr")
        is not None
    )
    assert date_.paragraph_format.space_before.pt == STYLES["date"].before
    assert subject.runs[0].bold
    assert signer.runs[0].bold and signer.text == "Nadia Contoh"
    assert all(
        p.paragraph_format.alignment != WD_ALIGN_PARAGRAPH.JUSTIFY for p in document.paragraphs
    )
    assert document.core_properties.author == "Nadia"


def test_pdf_is_one_a4_page_with_embedded_font_and_same_text_as_docx():
    reader = PdfReader(io.BytesIO(letter_to_pdf(LETTER, author="Nadia", contact=CONTACT)))
    document = Document(io.BytesIO(letter_to_docx(LETTER, author="Nadia", contact=CONTACT)))

    page = reader.pages[0]
    fonts = {font["/BaseFont"] for font in page["/Resources"]["/Font"].get_object().values()}
    assert len(reader.pages) == 1
    assert (round(float(page.mediabox.width)), round(float(page.mediabox.height))) == (595, 842)
    assert all("Caladea" in str(font) for font in fonts)
    assert reader.metadata.author == "Nadia"
    pdf_words = page.extract_text().split()
    docx_words = " ".join(p.text for p in document.paragraphs).split()
    assert pdf_words == docx_words


def test_long_letter_continues_on_next_pdf_page():
    text = "\n\n".join(["Paragraf isi yang cukup panjang. " * 12] * 12)

    assert len(PdfReader(io.BytesIO(letter_to_pdf(text))).pages) == 2


def test_letter_block_is_immutable():
    block = LetterBlock("body", ("a",))

    with pytest.raises(AttributeError):
        block.role = "name"  # type: ignore[misc]
